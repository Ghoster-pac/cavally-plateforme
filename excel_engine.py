"""
Moteur Excel pour la plateforme de suivi CAVALLY.

- parse_structure() : lit le classeur modele et decrit les lots, taches et lignes.
- read_cells()       : extrait les valeurs saisies (avancement / statut / observation).
- build_export()     : genere le classeur rempli en patchant directement le XML
                       du modele d'origine => le graphique, les styles, les mises
                       en forme conditionnelles et les listes deroulantes sont
                       conserves a l'identique.
"""
from __future__ import annotations

import io
import re
import zipfile
from xml.sax.saxutils import escape

import openpyxl

# ----------------------------------------------------------------------------
# Helpers colonnes
# ----------------------------------------------------------------------------


def col_to_num(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def num_to_col(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _slug(text: str) -> str:
    t = (text or "").strip().lower()
    t = (t.replace("é", "e").replace("è", "e").replace("ê", "e").replace("ë", "e")
          .replace("à", "a").replace("â", "a").replace("î", "i").replace("ï", "i")
          .replace("ô", "o").replace("û", "u").replace("ù", "u").replace("ç", "c")
          .replace("œ", "oe"))
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t or "tache"


# ----------------------------------------------------------------------------
# Lecture de la structure
# ----------------------------------------------------------------------------

SHEET_LABELS = {
    "AVANCEMENT PBS": "Plomberie (PBS)",
    "AVANCEMENT CVC": "CVCD (CVC)",
    "AVCT CFO": "CFO (Electricite)",
    "AVANCEMENT CFA-SSI": "CFA / SSI",
}

# ordre d'affichage souhaite
SHEET_ORDER = ["AVANCEMENT PBS", "AVANCEMENT CVC", "AVCT CFO", "AVANCEMENT CFA-SSI"]

LOT_KEY = {
    "AVANCEMENT PBS": "pbs",
    "AVANCEMENT CVC": "cvc",
    "AVCT CFO": "cfo",
    "AVANCEMENT CFA-SSI": "cfa",
}


def parse_structure(template_path: str) -> dict:
    wb = openpyxl.load_workbook(template_path, data_only=False)
    lots = []
    for sheet in SHEET_ORDER:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]

        # ligne contenant NIVEAU en colonne A
        nivel_row = None
        for r in range(1, min(ws.max_row, 30) + 1):
            if str(ws.cell(r, 1).value or "").strip().upper() == "NIVEAU":
                nivel_row = r
                break
        if nivel_row is None:
            continue

        # sous-en-tete : premiere ligne >= nivel_row ou la colonne D vaut AVANCEMENT
        sub_row = None
        for r in range(nivel_row, nivel_row + 4):
            if str(ws.cell(r, 4).value or "").strip().upper() == "AVANCEMENT":
                sub_row = r
                break
        if sub_row is None:
            sub_row = nivel_row + 1
        group_row = sub_row - 1
        data_start = sub_row + 1

        # taches : cellules de la ligne de groupe
        merged = {}
        for rng in ws.merged_cells.ranges:
            merged[(rng.min_row, rng.min_col)] = rng.max_col
        covered = set()
        for rng in ws.merged_cells.ranges:
            if rng.min_row == group_row:
                for c in range(rng.min_col, rng.max_col + 1):
                    covered.add(c)

        tasks = []
        seen_ids = set()
        for c in range(4, ws.max_column + 1):
            val = ws.cell(group_row, c).value
            if val is None or str(val).strip() == "":
                continue
            if c in covered and (group_row, c) not in merged:
                continue  # cellule couverte par une fusion, pas un debut
            label = str(val).strip()
            tid = _slug(label)
            base = tid
            i = 2
            while tid in seen_ids:
                tid = f"{base}-{i}"
                i += 1
            seen_ids.add(tid)
            tasks.append({"id": tid, "label": label, "col": c})

        # le dernier groupe est l'avancement global
        global_task = None
        if tasks and "GLOBAL" in tasks[-1]["label"].upper():
            global_task = tasks.pop()

        raw_rows = []
        summary_rows = []
        for r in range(data_start, ws.max_row + 1):
            a = ws.cell(r, 1).value
            b = ws.cell(r, 2).value
            a = str(a).strip() if a is not None else ""
            b = str(b).strip() if b is not None else ""
            if b.upper() == "MOYENNE":
                summary_rows.append((r, a))
                continue
            if not b:
                continue
            raw_rows.append((r, a, b))

        # bornes de bloc fournies par les lignes MOYENNE (le niveau y est explicite)
        bounds = []
        last_lvl = ""
        for sr, slvl in summary_rows:
            if slvl:
                last_lvl = slvl
            bounds.append((sr, last_lvl))
        bounds.sort()

        def infer_level(r):
            for sr, lvl in bounds:
                if r < sr and lvl:
                    return lvl
            return last_lvl

        rows = []
        for r, a, b in raw_rows:
            niveau = a if a else infer_level(r)
            rows.append({"row": r, "niveau": niveau, "designation": b})

        levels = []
        for row in rows:
            if row["niveau"] and row["niveau"] not in levels:
                levels.append(row["niveau"])

        lots.append({
            "key": LOT_KEY.get(sheet, _slug(sheet)),
            "sheet": sheet,
            "label": SHEET_LABELS.get(sheet, sheet),
            "nivel_row": nivel_row,
            "sub_row": sub_row,
            "group_row": group_row,
            "data_start": data_start,
            "niveau_col": 1,
            "desig_col": 2,
            "tasks": tasks,
            "global": {
                "label": global_task["label"] if global_task else "AVANCEMENT GLOBAL",
                "moyenne_col": global_task["col"] if global_task else None,
                "statut_col": (global_task["col"] + 1) if global_task else None,
                "commentaire_col": (global_task["col"] + 2) if global_task else None,
            },
            "rows": rows,
            "summary_rows": summary_rows,
            "levels": levels,
        })

    return {"project": "CAVALLY", "lots": lots}


# ----------------------------------------------------------------------------
# Lecture des valeurs saisies
# ----------------------------------------------------------------------------


def read_cells(template_path: str, structure: dict) -> dict:
    """Retourne {sheet: {(row, col): value}} pour les cellules de saisie."""
    wb = openpyxl.load_workbook(template_path, data_only=True)
    out = {}
    for lot in structure["lots"]:
        ws = wb[lot["sheet"]]
        sheet_vals = {}
        for row in lot["rows"]:
            r = row["row"]
            for t in lot["tasks"]:
                for off in (0, 1, 2):
                    v = ws.cell(r, t["col"] + off).value
                    if v is not None:
                        sheet_vals[(r, t["col"] + off)] = v
            if lot["global"]["commentaire_col"]:
                v = ws.cell(r, lot["global"]["commentaire_col"]).value
                if v is not None:
                    sheet_vals[(r, lot["global"]["commentaire_col"])] = v
        out[lot["sheet"]] = sheet_vals
    return out


# ----------------------------------------------------------------------------
# Construction du classeur exporte (patch XML)
# ----------------------------------------------------------------------------

ST_ORDER = ["Bloque", "Non conforme", "A verifier", "En cours", "Termine"]
ST_LABEL = {
    "Non demarre": "Non demarre",
    "En cours": "En cours",
    "Termine": "Termine",
    "Bloque": "Bloque",
    "Non conforme": "Non conforme",
    "A verifier": "A verifier",
}
# valeurs reelles des listes deroulantes du classeur (avec accents)
STATUS_VALUES = ["Non démarré", "En cours", "Terminé", "Bloqué", "Non conforme", "À vérifier"]


def _norm_status(s: str) -> str:
    if not s:
        return "Non démarré"
    return s


def _cell_xml(ref: str, style: str | None, kind: str, value) -> str:
    s = f' s="{style}"' if style else ""
    if kind == "f":
        return f'<c r="{ref}"{s}><f>{escape(str(value))}</f></c>'
    if kind == "n":
        return f'<c r="{ref}"{s}><v>{value}</v></c>'
    # texte inline (evite de toucher sharedStrings)
    return (f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">'
            f'{escape(str(value))}</t></is></c>')


def _match_cell(inner: str, ref: str):
    return re.search(r'<c r="%s"(?:\s[^>]*?)?(?:/>|>.*?</c>)' % re.escape(ref), inner, re.DOTALL)


def _strip_formula_caches(xml: str) -> str:
    """Supprime les valeurs en cache des cellules de formule pour forcer
    le recalcul a l'ouverture (Excel comme LibreOffice)."""
    def repl(m):
        cell = m.group(0)
        if "<f" not in cell:
            return cell
        cell = re.sub(r"<v>.*?</v>", "", cell, flags=re.DOTALL)
        return cell

    return re.sub(r"<c\b[^>]*?(?:/>|>.*?</c>)", repl, xml, flags=re.DOTALL)


def _set_formula(xml: str, ref: str, formula: str) -> str:
    m = _match_cell(xml, ref)
    if not m:
        return xml
    st = re.search(r'\bs="(\d+)"', m.group(0))
    style = f' s="{st.group(1)}"' if st else ""
    new = f'<c r="{ref}"{style}><f>{escape(formula)}</f></c>'
    return xml[:m.start()] + new + xml[m.end():]


def _set_text(xml: str, ref: str, text: str) -> str:
    m = _match_cell(xml, ref)
    if not m:
        return xml
    st = re.search(r'\bs="(\d+)"', m.group(0))
    style = f' s="{st.group(1)}"' if st else ""
    new = (f'<c r="{ref}"{style} t="inlineStr"><is><t xml:space="preserve">'
           f'{escape(text)}</t></is></c>')
    return xml[:m.start()] + new + xml[m.end():]


def _ranges(lot: dict):
    start = lot["data_start"]
    end = max((row["row"] for row in lot["rows"]), default=start)
    sheet = lot["sheet"]
    col = num_to_col(lot["global"]["moyenne_col"])
    return (f"'{sheet}'!$B${start}:$B${end}", f"'{sheet}'!${col}${start}:${col}${end}")


def _set_cell(inner: str, ref: str, col: int, style: str | None, kind: str, value) -> str:
    new_cell = _cell_xml(ref, style, kind, value)
    m = _match_cell(inner, ref)
    if m:
        st = re.search(r'\bs="(\d+)"', m.group(0))
        if st:
            new_cell = _cell_xml(ref, st.group(1), kind, value)
        return inner[:m.start()] + new_cell + inner[m.end():]
    # insertion au bon endroit
    cells = list(re.finditer(r'<c\b[^>]*?(?:/>|>.*?</c>)', inner, re.DOTALL))
    for cm in cells:
        r = re.search(r'r="([A-Z]+)\d+"', cm.group(0))
        if r and col_to_num(r.group(1)) > col:
            return inner[:cm.start()] + new_cell + inner[cm.start():]
    return inner + new_cell


def _sheet_cell_updates_for_lot(lot: dict, data: dict) -> dict:
    """{row: {col: (kind, value)}} pour une feuille de lot."""
    updates = {}
    task_by_id = {t["id"]: t for t in lot["tasks"]}
    g = lot["global"]
    for row in lot["rows"]:
        r = row["row"]
        d = data.get(r) or {}
        av = d.get("av", {})
        st = d.get("st", {})
        obs = d.get("obs", {})
        cells = {}
        if row.get("niveau"):
            cells[lot["niveau_col"]] = ("s", row["niveau"])
        for t in lot["tasks"]:
            c = t["col"]
            a = av.get(t["id"])
            if a is not None and a != "":
                cells[c] = ("n", _fmt_num(a))
            sv = st.get(t["id"])
            if sv:
                cells[c + 1] = ("n" if False else "s", _norm_status(sv))
            ov = obs.get(t["id"])
            if ov not in (None, ""):
                cells[c + 2] = ("s", ov)
        # moyenne (formule)
        if g["moyenne_col"]:
            av_cols = [num_to_col(t["col"]) + str(r) for t in lot["tasks"]]
            if av_cols:
                cells[g["moyenne_col"]] = ("f", "IFERROR(AVERAGE(%s),0)" % ",".join(av_cols))
            # statut global (formule basee sur les statuts de taches)
            if g["statut_col"]:
                st_refs = [num_to_col(t["col"] + 1) + str(r) for t in lot["tasks"]]
                cells[g["statut_col"]] = ("f", _status_formula(st_refs, g["moyenne_col"], r))
            if g["commentaire_col"]:
                cv = d.get("comment")
                if cv not in (None, ""):
                    cells[g["commentaire_col"]] = ("s", cv)
        updates[r] = cells
    return updates


def _status_formula(st_refs, moyenne_col, r):
    ors = lambda op: ",".join(f'{ref}="{op}"' for ref in st_refs)
    moy = f"{num_to_col(moyenne_col)}{r}" if moyenne_col else "0"
    return (
        f'IF(OR({ors("Bloqué")}),"Bloqué",'
        f'IF(OR({ors("Non conforme")}),"Non conforme",'
        f'IF(OR({ors("À vérifier")}),"À vérifier",'
        f'IF({moy}>=1,"Terminé",'
        f'IF({moy}>0,"En cours","Non démarré")))))'
    )


def _fmt_num(v) -> str:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return "0"
    if f == int(f):
        return str(int(f))
    return repr(f)


def _patch_sheet_rows(xml: str, updates: dict) -> str:
    if not updates:
        return xml
    # style de reference par colonne (pour les cellules vides inserees)
    style_by_col = {}
    for m in re.finditer(r'<c r="([A-Z]+)\d+"[^>]*?\bs="(\d+)"', xml):
        style_by_col.setdefault(m.group(1), m.group(2))
    pattern = re.compile(r'(<row\b[^>]*>)(.*?)(</row>)', re.DOTALL)

    def repl(m):
        open_tag, inner, close = m.group(1), m.group(2), m.group(3)
        rm = re.search(r'r="(\d+)"', open_tag)
        if not rm:
            return m.group(0)
        r = int(rm.group(1))
        cells = updates.get(r)
        if not cells:
            return m.group(0)
        for col in sorted(cells):
            kind, value = cells[col]
            ref = f"{num_to_col(col)}{r}"
            fallback = style_by_col.get(num_to_col(col))
            inner = _set_cell(inner, ref, col, fallback, kind, value)
        return open_tag + inner + close

    return pattern.sub(repl, xml)


def build_export(template_path: str, structure: dict, data_by_sheet: dict) -> bytes:
    """Genere le xlsx rempli a partir du modele, en patchant le XML."""
    zin = zipfile.ZipFile(template_path, "r")
    names = zin.namelist()

    # correspondance feuille -> fichier xml
    wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
    rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_map = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    sheet_file = {}
    for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="([^"]+)"', wb_xml):
        name, rid = m.group(1), m.group(2)
        target = rel_map.get(rid, "")
        target = target.replace("../", "")
        if not target.startswith("xl/"):
            target = "xl/" + target.lstrip("/")
        sheet_file[name] = target

    patched = {}
    for lot in structure["lots"]:
        sheet = lot["sheet"]
        path = sheet_file.get(sheet)
        if not path:
            continue
        updates = _sheet_cell_updates_for_lot(lot, data_by_sheet.get(sheet, {}))
        xml = zin.read(path).decode("utf-8")
        xml = _patch_sheet_rows(xml, updates)
        patched[path] = xml

    # corriger les formules a liens externes -> references internes
    for sheet in ("TABLEAU DE BORD", "AVANCEMENT PAR NIVEAU", "RECAP GENERAL"):
        path = sheet_file.get(sheet)
        if path and path not in patched:
            patched[path] = zin.read(path).decode("utf-8")

    if "TABLEAU DE BORD" in sheet_file:
        p = sheet_file["TABLEAU DE BORD"]
        xml = patched[p]
        lots_by_sheet = {l["sheet"]: l for l in structure["lots"]}
        row_sheet = {4: "AVANCEMENT CVC", 5: "AVANCEMENT PBS", 6: "AVCT CFO",
                     7: "AVANCEMENT CFA-SSI", 8: "AVANCEMENT CFA-SSI"}
        for r, sheet in row_sheet.items():
            lot = lots_by_sheet.get(sheet)
            if not lot:
                continue
            brange, avrange = _ranges(lot)
            cnt = f'COUNTIF({brange},"<>")-COUNTIF({brange},"MOYENNE")'
            avg = (f'IFERROR(SUMPRODUCT(({brange}<>"MOYENNE")*{avrange})'
                   f'/SUMPRODUCT(--({brange}<>"MOYENNE")),0)')
            xml = _set_formula(xml, f"B{r}", cnt)
            xml = _set_formula(xml, f"C{r}", avg)
        xml = _set_formula(xml, "B10", "IFERROR('RECAP GENERAL'!B11,0)")
        patched[p] = xml

    if "RECAP GENERAL" in sheet_file:
        p = sheet_file["RECAP GENERAL"]
        xml = patched[p]
        lots_by_sheet = {l["sheet"]: l for l in structure["lots"]}
        recap_sheet = {6: "AVANCEMENT PBS", 7: "AVANCEMENT CVC", 8: "AVCT CFO",
                       9: "AVANCEMENT CFA-SSI"}
        for r, sheet in recap_sheet.items():
            lot = lots_by_sheet.get(sheet)
            if not lot:
                continue
            brange, avrange = _ranges(lot)
            avg = (f'IFERROR(SUMPRODUCT(({brange}<>"MOYENNE")*{avrange})'
                   f'/SUMPRODUCT(--({brange}<>"MOYENNE")),0)')
            xml = _set_formula(xml, f"B{r}", avg)
        patched[p] = xml

    if "AVANCEMENT PAR NIVEAU" in sheet_file:
        p = sheet_file["AVANCEMENT PAR NIVEAU"]
        xml = patched[p]
        lot_by_col = {"B": "AVANCEMENT CVC", "C": "AVANCEMENT PBS",
                      "D": "AVCT CFO", "E": "AVANCEMENT CFA-SSI",
                      "F": "AVANCEMENT CFA-SSI"}
        glob_col = {l["sheet"]: l["global"]["moyenne_col"] for l in structure["lots"]}
        # correspondance libelle de zone -> niveau reel des feuilles de lot
        niveau_par_ligne = {2: "SS-2", 3: "SS-1", 4: "RDC", 5: "R+1", 6: "R+2",
                            7: "R+3", 8: "R+4", 9: "R+5", 10: "R+6", 11: "R+6",
                            12: "__AUCUN__"}
        for rr in range(2, 13):
            niveau = niveau_par_ligne.get(rr)
            for col, sheet in lot_by_col.items():
                gc = glob_col.get(sheet)
                if not gc:
                    continue
                g = num_to_col(gc)
                f = (f"IFERROR(AVERAGEIFS('{sheet}'!${g}:${g},'{sheet}'!$A:$A,"
                     f"\"{niveau}\",'{sheet}'!$B:$B,\"<>MOYENNE\"),0)")
                ref = f"{col}{rr}"
                m = _match_cell(xml, ref)
                if m:
                    xml = xml[:m.start()] + f'<c r="{ref}"><f>{escape(f)}</f></c>' + xml[m.end():]
        # note sur les zones non cartographiees
        xml = _set_text(xml, "G12",
                        "Locaux techniques : repartis sur plusieurs niveaux")
        patched[p] = xml

    # ecriture du nouvel archive
    drop_parts = set()
    for n in names:
        if n.startswith("xl/externalLinks/"):
            drop_parts.add(n)
        if n == "xl/calcChain.xml":
            drop_parts.add(n)

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n in drop_parts:
                continue
            if n in patched:
                content = patched[n].encode("utf-8")
            else:
                content = zin.read(n)
            if n.startswith("xl/worksheets/") and n.endswith(".xml"):
                s = content.decode("utf-8")
                content = _strip_formula_caches(s).encode("utf-8")
            if n == "xl/workbook.xml":
                s = content.decode("utf-8")
                s = re.sub(r"<externalReferences>.*?</externalReferences>", "", s, flags=re.DOTALL)
                s = s.replace('forceFullCalc="1"', 'forceFullCalc="1" fullCalcOnLoad="1"')
                content = s.encode("utf-8")
            elif n == "xl/_rels/workbook.xml.rels":
                s = content.decode("utf-8")
                s = re.sub(r'<Relationship[^>]*externalLink[^>]*/>', "", s)
                s = re.sub(r'<Relationship[^>]*calcChain[^>]*/>', "", s)
                content = s.encode("utf-8")
            elif n == "[Content_Types].xml":
                s = content.decode("utf-8")
                s = re.sub(r'<Override[^>]*externalLinks/[^>]*?/>', "", s)
                s = re.sub(r'<Override[^>]*calcChain\.xml[^>]*?/>', "", s)
                content = s.encode("utf-8")
            zout.writestr(n, content)
    zin.close()
    return out.getvalue()


# ----------------------------------------------------------------------------
# Import d'un classeur rempli
# ----------------------------------------------------------------------------


def parse_import(file_path: str, structure: dict) -> dict:
    """Lit un xlsx rempli et retourne {sheet: {(row,col): value}} des saisies."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    out = {}
    for lot in structure["lots"]:
        if lot["sheet"] not in wb.sheetnames:
            continue
        ws = wb[lot["sheet"]]
        vals = {}
        for row in lot["rows"]:
            r = row["row"]
            for t in lot["tasks"]:
                for off in (0, 1, 2):
                    v = ws.cell(r, t["col"] + off).value
                    if v is not None:
                        vals[(r, t["col"] + off)] = v
            gc = lot["global"]["commentaire_col"]
            if gc:
                v = ws.cell(r, gc).value
                if v is not None:
                    vals[(r, gc)] = v
        out[lot["sheet"]] = vals
    return out


if __name__ == "__main__":
    import json
    import sys

    p = sys.argv[1] if len(sys.argv) > 1 else "files/cavally_template.xlsx"
    st = parse_structure(p)
    print(json.dumps({k: (v if k != "lots" else [
        {kk: (vv if kk not in ("rows",) else len(vv)) for kk, vv in l.items()} for l in v
    ]) for k, v in st.items()}, ensure_ascii=False, indent=2))
